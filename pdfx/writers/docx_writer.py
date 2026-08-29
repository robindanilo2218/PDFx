"""Salida Word (.docx) con python-docx."""

from __future__ import annotations

import logging
from pathlib import Path

from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Emu, Inches, Pt, RGBColor

from ..config import Settings
from ..model import Block, Document, ImageRef, Table

log = logging.getLogger(__name__)

_MAX_TABLE_COLS = 60      # Word se vuelve inestable con rejillas enormes


def _content_width(section) -> int:
    return int(section.page_width - section.left_margin - section.right_margin)


def _add_table(docx, table: Table) -> None:
    rows = table.normalized()
    if not rows:
        return
    n_cols = min(len(rows[0]), _MAX_TABLE_COLS)
    if n_cols < 1:
        return

    t = docx.add_table(rows=0, cols=n_cols)
    try:
        t.style = "Table Grid"
    except KeyError:
        pass

    for r_idx, row in enumerate(rows):
        cells = t.add_row().cells
        for c_idx in range(n_cols):
            text = row[c_idx] if c_idx < len(row) else ""
            cell = cells[c_idx]
            cell.text = ""
            para = cell.paragraphs[0]
            run = para.add_run(text)
            if r_idx == 0 and table.has_header:
                run.bold = True
    docx.add_paragraph()


def _add_image(docx, ref: ImageRef, max_width_emu: int) -> None:
    if not ref.abs_path or not ref.abs_path.exists():
        return
    try:
        # 96 ppp es la referencia de Word para pixeles.
        natural = Emu(int(Inches(ref.width / 96.0)))
        width = min(natural, max_width_emu) if ref.width else max_width_emu
        docx.add_picture(str(ref.abs_path), width=width)
        docx.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if ref.caption:
            cap = docx.add_paragraph(ref.caption)
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.runs[0].italic = True
            cap.runs[0].font.size = Pt(9)
    except Exception as exc:
        log.warning("No se pudo insertar %s en Word: %s", ref.abs_path, exc)
        docx.add_paragraph(f"[Imagen no insertada: {ref.abs_path.name}]")


def _add_block(docx, block: Block) -> None:
    text = block.text.strip()
    if not text:
        return

    if block.kind == "heading":
        level = max(1, min(6, block.level or 2))
        try:
            docx.add_heading(text, level=level)
        except KeyError:
            para = docx.add_paragraph()
            run = para.add_run(text)
            run.bold = True
            run.font.size = Pt(max(11, 20 - 2 * level))
        return

    if block.kind == "list_item":
        body = text
        if block.marker and body.startswith(block.marker):
            body = body[len(block.marker):].lstrip()
        style = "List Number" if block.ordered else "List Bullet"
        try:
            para = docx.add_paragraph(body, style=style)
        except KeyError:
            para = docx.add_paragraph(("- " if not block.ordered else "1. ") + body)
        if block.indent:
            para.paragraph_format.left_indent = Inches(0.25 * min(block.indent, 4))
        return

    if block.kind == "caption":
        para = docx.add_paragraph(text)
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.runs[0]
        run.italic = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        return

    docx.add_paragraph(text)


def write(doc: Document, dest: Path, settings: Settings,
          redaction_note: str = "") -> Path:
    docx = DocxDocument()

    docx.core_properties.title = doc.title or doc.source_path.stem
    docx.core_properties.author = doc.author or "PDFx"
    docx.core_properties.comments = (
        f"Convertido localmente desde {doc.source_path.name}"
        + (f" | {redaction_note}" if redaction_note else "")
    )

    section = docx.sections[0]
    max_width = _content_width(section)

    docx.add_heading(doc.title.strip() or doc.source_path.stem, level=0)
    if redaction_note:
        note = docx.add_paragraph(redaction_note)
        note.runs[0].italic = True
        note.runs[0].font.size = Pt(9)

    for p_idx, page in enumerate(doc.pages):
        if p_idx > 0 and settings.docx_page_breaks:
            docx.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        for el in page.elements:
            if isinstance(el, Block):
                _add_block(docx, el)
            elif isinstance(el, Table):
                _add_table(docx, el)
            elif isinstance(el, ImageRef):
                _add_image(docx, el, max_width)

    if doc.warnings:
        docx.add_page_break()
        docx.add_heading("Avisos de la conversion", level=2)
        for w in doc.warnings:
            docx.add_paragraph(w, style="List Bullet")

    dest.parent.mkdir(parents=True, exist_ok=True)
    docx.save(str(dest))
    return dest
