"""Salida Excel (.xlsx) con openpyxl.

Modo 'tables': una hoja por tabla detectada, mas una hoja indice.
Modo 'full':   ademas una hoja con todo el texto, pagina a pagina.
"""

from __future__ import annotations

import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..config import Settings
from ..model import Document

_HEADER_FILL = PatternFill("solid", fgColor="DDE5EF")
_TITLE_FONT = Font(bold=True, size=12)
_HEADER_FONT = Font(bold=True)
_INVALID_SHEET = re.compile(r"[\\/*?:\[\]]")
_MAX_COL_WIDTH = 60
_EXCEL_MAX_CELL = 32767


def _sheet_name(base: str, used: set[str]) -> str:
    name = _INVALID_SHEET.sub("-", base).strip()[:28] or "Hoja"
    candidate, n = name, 2
    while candidate.lower() in used:
        suffix = f"_{n}"
        candidate = name[: 31 - len(suffix)] + suffix
        n += 1
    used.add(candidate.lower())
    return candidate


def _coerce(value: str):
    """Convierte a numero cuando es claramente numerico (formato ES o EN)."""
    text = (value or "").strip()
    if not text or len(text) > 40:
        return text[:_EXCEL_MAX_CELL]
    cleaned = text.replace("%", "").replace("€", "").replace("$", "").strip()
    if not re.fullmatch(r"[-+]?[\d.,]+", cleaned):
        return text
    # 1.234,56 (ES)  vs  1,234.56 (EN)
    if "," in cleaned and "." in cleaned:
        if cleaned.rfind(",") > cleaned.rfind("."):
            cleaned = cleaned.replace(".", "").replace(",", ".")
        else:
            cleaned = cleaned.replace(",", "")
    elif "," in cleaned:
        cleaned = cleaned.replace(",", ".") if cleaned.count(",") == 1 else cleaned.replace(",", "")
    try:
        num = float(cleaned)
    except ValueError:
        return text
    return int(num) if num.is_integer() and abs(num) < 1e15 else num


def _autosize(ws, rows: list[list], max_scan: int = 200) -> None:
    if not rows:
        return
    n_cols = max(len(r) for r in rows)
    for c in range(n_cols):
        widest = 8
        for row in rows[:max_scan]:
            if c < len(row):
                widest = max(widest, min(len(str(row[c] or "")), _MAX_COL_WIDTH))
        ws.column_dimensions[get_column_letter(c + 1)].width = widest + 2


def write(doc: Document, dest: Path, settings: Settings,
          redaction_note: str = "") -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()

    index = wb.create_sheet(_sheet_name("Indice", used))
    index["A1"] = doc.title or doc.source_path.stem
    index["A1"].font = _TITLE_FONT
    index["A2"] = f"Origen: {doc.source_path.name}"
    index["A3"] = (
        f"{doc.n_pages} paginas | {doc.n_tables} tablas | {doc.n_images} imagenes"
        f" | OCR en {doc.n_ocr_pages} paginas"
    )
    if redaction_note:
        index["A4"] = redaction_note
        index["A4"].font = Font(italic=True)

    head_row = 6
    for col, title in enumerate(("Hoja", "Pagina", "Filas", "Columnas",
                                 "Deteccion", "Fiabilidad"), start=1):
        cell = index.cell(row=head_row, column=col, value=title)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL

    row_ptr = head_row + 1
    n_tables = 0

    for page in doc.pages:
        for t_idx, table in enumerate(page.tables, start=1):
            rows = table.normalized()
            if not rows:
                continue
            n_tables += 1
            name = _sheet_name(f"p{page.number}_t{t_idx}", used)
            ws = wb.create_sheet(name)
            for r_idx, row in enumerate(rows, start=1):
                for c_idx, value in enumerate(row, start=1):
                    cell = ws.cell(row=r_idx, column=c_idx, value=_coerce(value))
                    if r_idx == 1 and table.has_header:
                        cell.font = _HEADER_FONT
                        cell.fill = _HEADER_FILL
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
            if table.has_header and len(rows) > 1:
                ws.freeze_panes = "A2"
                ws.auto_filter.ref = (
                    f"A1:{get_column_letter(len(rows[0]))}{len(rows)}"
                )
            _autosize(ws, rows)

            index.cell(row=row_ptr, column=1, value=name).hyperlink = f"#'{name}'!A1"
            index.cell(row=row_ptr, column=1).style = "Hyperlink"
            index.cell(row=row_ptr, column=2, value=page.number)
            index.cell(row=row_ptr, column=3, value=table.n_rows)
            index.cell(row=row_ptr, column=4, value=table.n_cols)
            index.cell(row=row_ptr, column=5, value=table.source)
            index.cell(row=row_ptr, column=6, value=round(table.confidence, 2))
            row_ptr += 1

    if n_tables == 0:
        index.cell(row=row_ptr, column=1,
                   value="No se detecto ninguna tabla en el documento.")

    if settings.xlsx_mode == "full" or n_tables == 0:
        ws = wb.create_sheet(_sheet_name("Texto", used))
        headers = ("Pagina", "Tipo", "Nivel", "Texto")
        for c, title in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=c, value=title)
            cell.font = _HEADER_FONT
            cell.fill = _HEADER_FILL
        r = 2
        for page in doc.pages:
            for block in page.blocks:
                text = block.text.strip()
                if not text:
                    continue
                ws.cell(row=r, column=1, value=page.number)
                ws.cell(row=r, column=2, value=block.kind)
                ws.cell(row=r, column=3, value=block.level or "")
                cell = ws.cell(row=r, column=4, value=text[:_EXCEL_MAX_CELL])
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                r += 1
        ws.freeze_panes = "A2"
        for col, width in zip("ABCD", (9, 12, 7, 110)):
            ws.column_dimensions[col].width = width

    _autosize(index, [["Hoja", "Pagina", "Filas", "Columnas", "Deteccion", "Fiabilidad"]])
    index.column_dimensions["A"].width = 22
    index.column_dimensions["E"].width = 14

    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(dest))
    return dest
