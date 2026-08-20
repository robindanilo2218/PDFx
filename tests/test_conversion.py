"""Conversion de punta a punta sobre los PDF de prueba."""

import re

import pytest

from pdfx.config import RedactionSettings, Settings
from pdfx.convert import convert_file

FILAS_ESPERADAS = {
    "SIE-4520B": ("Variador de frecuencia 7,5 kW", "2", "1.240,50"),
    "ABB-7710": ("Contactor tripolar 32 A", "12", "86,20"),
    "FES-0912": ("Cilindro neumatico 63x200", "4", "312,00"),
    "SKF-6205": ("Rodamiento rigido de bolas", "24", "18,75"),
}


def convertir(src, tmp_path, **kw):
    settings = Settings(out_dir=tmp_path, **kw)
    result = convert_file(src, settings)
    assert not result.errors, result.errors
    return result


# --------------------------------------------------------------- nativo ---
def test_nativo_estructura_markdown(nativo, tmp_path):
    md = convertir(nativo, tmp_path, formats=["md"]).outputs["md"].read_text()
    # El pie de avisos tambien usa vinetas ("- Tesseract no encontrado: ..."),
    # asi que se descarta: aqui se comprueba la estructura del documento, no
    # si esta maquina tiene OCR instalado.
    md = md.split("**Avisos de la conversion**")[0]

    # El titulo del documento no se repite con el primer encabezado.
    assert md.count("# Informe de mantenimiento preventivo") == 1
    # Los apartados numerados quedan en nivel 2, bajo el titulo.
    for n in range(1, 6):
        assert re.search(rf"^## {n}\. ", md, re.M), f"falta el apartado {n}"
    # Las cuatro vinetas son items de lista, no un parrafo corrido.
    assert md.count("\n- ") == 4
    assert "Figura 1" in md and "*Figura 1" in md   # pie de imagen en cursiva


def test_nativo_tabla_completa(nativo, tmp_path):
    doc = convertir(nativo, tmp_path, formats=["md"]).document
    tablas = [t for p in doc.pages for t in p.tables]
    assert len(tablas) == 1
    tabla = tablas[0]
    assert tabla.rows[0] == ["Referencia", "Descripcion", "Cantidad", "Precio (EUR)"]
    for fila in tabla.rows[1:]:
        ref = fila[0]
        assert tuple(fila[1:]) == FILAS_ESPERADAS[ref]


def test_nativo_imagen_extraida(nativo, tmp_path):
    result = convertir(nativo, tmp_path, formats=["md"])
    imagenes = [i for p in result.document.pages for i in p.images]
    assert len(imagenes) == 1
    assert imagenes[0].abs_path.exists()
    assert imagenes[0].abs_path.stat().st_size > 1000


def test_excel_convierte_numeros_en_formato_espanol(nativo, tmp_path):
    from openpyxl import load_workbook
    salida = convertir(nativo, tmp_path, formats=["xlsx"]).outputs["xlsx"]
    wb = load_workbook(salida)
    assert "Indice" in wb.sheetnames
    hoja = wb[wb.sheetnames[1]]
    filas = list(hoja.iter_rows(values_only=True))
    assert filas[1][2] == 2            # cantidad como entero
    assert filas[1][3] == 1240.5       # "1.240,50" -> numero


def test_word_usa_estilos_reales(nativo, tmp_path):
    from docx import Document as Docx
    salida = convertir(nativo, tmp_path, formats=["docx"]).outputs["docx"]
    docx = Docx(str(salida))
    estilos = {p.style.name for p in docx.paragraphs if p.text.strip()}
    assert "Heading 2" in estilos and "List Bullet" in estilos
    assert len(docx.tables) == 1
    assert len(docx.inline_shapes) == 1


def test_html_incrusta_las_imagenes(nativo, tmp_path):
    html = convertir(nativo, tmp_path, formats=["html"]).outputs["html"].read_text()
    assert "data:image/png;base64," in html
    assert "<table>" in html


# ---------------------------------------------------------- dos columnas ---
def test_dos_columnas_orden_de_lectura(dos_columnas, tmp_path):
    result = convertir(dos_columnas, tmp_path, formats=["md"])
    md = result.outputs["md"].read_text()
    assert re.findall(r"Bloque (\d)", md) == [str(i) for i in range(1, 9)]
    assert any("2 columnas" in n for p in result.document.pages for n in p.notes)


def test_columnas_desactivables(dos_columnas, tmp_path):
    result = convertir(dos_columnas, tmp_path, formats=["md"], detect_columns=False)
    assert not any("columnas" in n for p in result.document.pages for n in p.notes)


# -------------------------------------------------------------- opciones ---
def test_rango_de_paginas(nativo, tmp_path):
    doc = convertir(nativo, tmp_path, formats=["md"], page_range="2").document
    assert doc.n_pages == 1
    assert doc.pages[0].number == 2


def test_modo_rapido_deja_el_md_junto_al_pdf(nativo, tmp_path):
    import shutil
    copia = tmp_path / "copia.pdf"
    shutil.copy(nativo, copia)
    result = convert_file(copia, Settings(formats=["md"], fast_mode=True,
                                          extract_images=False))
    assert result.outputs["md"] == tmp_path / "copia.md"
    assert not (tmp_path / "copia_media").exists()


def test_sin_tablas(nativo, tmp_path):
    doc = convertir(nativo, tmp_path, formats=["md"], detect_tables=False).document
    assert doc.n_tables == 0


# --------------------------------------------------------------- saneado ---
def test_saneado_sin_fugas_en_ningun_formato(nativo, tmp_path):
    redaction = RedactionSettings(
        enabled=True,
        terms=["Talleres Munoz S.L.", "Javier Ortega"],
        detectors=["email", "phone", "serial"],
        style="label",
    )
    result = convertir(nativo, tmp_path, formats=["md", "html", "txt", "json"],
                       redaction=redaction)
    assert result.redaction.total > 0
    assert result.redaction.leaked == []

    secretos = ("Talleres Munoz", "Javier Ortega", "soporte@talleresmunoz.es",
                "912 445 780", "SIE-4520B")
    for fmt in ("md", "html", "txt", "json"):
        texto = result.outputs[fmt].read_text()
        for secreto in secretos:
            assert secreto not in texto, f"{secreto} sigue en {fmt}"

    informe = result.outputs["informe_saneado"].read_text()
    assert "Talleres Munoz S.L." in informe      # el informe si los lista
    assert "ningun termino marcado quedo" in informe


def test_saneado_puede_excluir_las_imagenes(nativo, tmp_path):
    redaction = RedactionSettings(enabled=True, terms=["ACME"], drop_images=True)
    result = convertir(nativo, tmp_path, formats=["md"], redaction=redaction)
    assert result.document.n_images == 0
    assert not (tmp_path / "media").exists()


# ------------------------------------------------------------------ OCR ---
@pytest.mark.slow
def test_escaneado_ocr_completo(escaneado, tmp_path, tesseract_disponible):
    if not tesseract_disponible:
        pytest.skip("Tesseract no esta instalado")
    result = convertir(escaneado, tmp_path, formats=["md"], ocr="auto")
    doc = result.document
    assert doc.n_ocr_pages == doc.n_pages
    assert all(p.ocr_conf > 80 for p in doc.pages)

    md = result.outputs["md"].read_text()
    assert "Informe de mantenimiento preventivo" in md
    assert md.count("\n- ") == 4                      # vinetas leidas por OCR
    assert any("enderezado" in n for p in doc.pages for n in p.notes)


@pytest.mark.slow
def test_escaneado_tabla_identica_a_la_nativa(escaneado, tmp_path, tesseract_disponible):
    if not tesseract_disponible:
        pytest.skip("Tesseract no esta instalado")
    doc = convertir(escaneado, tmp_path, formats=["md"], ocr="auto").document
    tablas = [t for p in doc.pages for t in p.tables]
    assert len(tablas) == 1, "el diagrama no debe confundirse con una tabla"
    tabla = tablas[0]
    assert tabla.rows[0] == ["Referencia", "Descripcion", "Cantidad", "Precio (EUR)"]
    for fila in tabla.rows[1:]:
        assert tuple(fila[1:]) == FILAS_ESPERADAS[fila[0]]


@pytest.mark.slow
def test_escaneado_recorta_el_esquema(escaneado, tmp_path, tesseract_disponible):
    if not tesseract_disponible:
        pytest.skip("Tesseract no esta instalado")
    doc = convertir(escaneado, tmp_path, formats=["md"], ocr="auto").document
    imagenes = [i for p in doc.pages for i in p.images]
    assert len(imagenes) == 1
    assert imagenes[0].kind == "region"
    assert imagenes[0].width > 200


def test_ocr_never_no_inventa_texto(escaneado, tmp_path):
    doc = convertir(escaneado, tmp_path, formats=["md"], ocr="never").document
    assert doc.char_count == 0
